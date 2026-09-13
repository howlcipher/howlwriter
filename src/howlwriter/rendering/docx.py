"""Native DOCX document renderer with APA 7 and academic formatting support."""

from __future__ import annotations

import html
import io
import re
from typing import Any

from howlwriter.domain.document import Document
from howlwriter.rendering.base import DocumentRenderer, RenderContext, RenderedArtifact


def _set_run_font(run: Any, ctx: RenderContext, *, size_pt: float | None = None) -> None:
    """Applies the requested font family and optional size to a run."""
    from docx.shared import Pt

    run.font.name = ctx.font_family
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _add_hyperlink(paragraph: Any, url: str, text: str) -> None:
    """Adds a clickable hyperlink run to a python-docx paragraph."""
    try:
        import docx.opc.constants
        from docx.oxml import parse_xml

        part = paragraph.part
        r_id = part.relate_to(
            url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        safe_text = html.escape(text)
        hyperlink_xml = (
            f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            f'r:id="{r_id}">'
            f'<w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr>'
            f'<w:t>{safe_text}</w:t></w:r></w:hyperlink>'
        )
        hyperlink = parse_xml(hyperlink_xml)
        paragraph._p.append(hyperlink)
    except Exception:
        # Fallback to plain text if XML manipulation fails
        run = paragraph.add_run(f"{text} ({url})")
        run.font.underline = True


def _add_inline_formatted(paragraph: Any, text: str, ctx: RenderContext) -> None:
    """Parses basic markdown formatting (**bold**, *italic*, `code`, [anchor](url)) into paragraph runs."""
    # Pattern to match links, bold, italic, and code
    pattern = re.compile(
        r"(\[(?P<link_text>[^\]]+)\]\((?P<link_url>[^\)]+)\))|"
        r"(\*\*(?P<bold>[^*]+)\*\*)|"
        r"(\*(?P<italic>[^*]+)\*)|"
        r"(`(?P<code>[^`]+)`)"
    )

    last_idx = 0
    for match in pattern.finditer(text):
        start, end = match.span()
        if start > last_idx:
            _set_run_font(paragraph.add_run(text[last_idx:start]), ctx)

        if match.group("link_text"):
            _add_hyperlink(
                paragraph, match.group("link_url"), match.group("link_text")
            )
        elif match.group("bold"):
            run = paragraph.add_run(match.group("bold"))
            run.bold = True
            _set_run_font(run, ctx)
        elif match.group("italic"):
            run = paragraph.add_run(match.group("italic"))
            run.italic = True
            _set_run_font(run, ctx)
        elif match.group("code"):
            run = paragraph.add_run(match.group("code"))
            run.font.name = "Courier New"

        last_idx = end

    if last_idx < len(text):
        _set_run_font(paragraph.add_run(text[last_idx:]), ctx)


def _add_page_number_header(section: Any, ctx: RenderContext) -> None:
    """Puts a right-aligned PAGE field in the section header (APA 7 placement)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    header_para = section.header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_para.add_run()
    run.font.name = ctx.font_family
    run.font.size = Pt(ctx.font_size_pt)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _add_title_page(doc: Any, title: str, lines: list[str], ctx: RenderContext) -> None:
    """Writes an APA-style title page: title in the upper half, then the
    author/affiliation/course lines, then a page break."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt

    for _ in range(6):
        doc.add_paragraph()
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.bold = True
    title_run.font.name = ctx.font_family
    title_run.font.size = Pt(ctx.font_size_pt)
    title_para.paragraph_format.space_after = Pt(ctx.font_size_pt)
    for line in lines:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(line)
        run.font.name = ctx.font_family
        run.font.size = Pt(ctx.font_size_pt)
    break_para = doc.add_paragraph()
    break_para.add_run().add_break(WD_BREAK.PAGE)


def _parse_markdown_table(lines: list[str]) -> tuple[list[str], list[list[str]]] | None:
    """Parses a markdown table into header and data rows."""
    if len(lines) < 2:
        return None

    # Check separator line (e.g. |---|---|)
    sep_line = lines[1].strip()
    if not re.match(r"^\|?[\s\-:|]+\|?$", sep_line):
        return None

    def split_row(line: str) -> list[str]:
        raw_cells = line.strip().split("|")
        # Remove empty ends from leading/trailing pipes
        if raw_cells and raw_cells[0].strip() == "":
            raw_cells = raw_cells[1:]
        if raw_cells and raw_cells[-1].strip() == "":
            raw_cells = raw_cells[:-1]
        return [c.strip() for c in raw_cells]

    headers = split_row(lines[0])
    if not headers:
        return None

    rows: list[list[str]] = []
    for line in lines[2:]:
        if "|" in line:
            cells = split_row(line)
            if cells:
                # Pad or truncate to match header length
                while len(cells) < len(headers):
                    cells.append("")
                rows.append(cells[: len(headers)])

    return headers, rows


class DocxRenderer(DocumentRenderer):
    """Renders a Document into a polished DOCX file with APA 7 layout conventions."""

    def render(
        self, document: Document, context: RenderContext | None = None
    ) -> RenderedArtifact:
        try:
            import docx
            from docx.shared import Inches, Pt, RGBColor
        except ImportError as exc:
            raise RuntimeError(
                "DOCX rendering requires 'python-docx'. "
                "Install it with: pip install 'howlwriter[docx]'"
            ) from exc

        ctx = context or RenderContext()
        title = ctx.title or document.title or "Untitled Document"

        doc = docx.Document()

        # Set margins from the formatting context
        for section in doc.sections:
            section.top_margin = Inches(ctx.margin_top_in)
            section.bottom_margin = Inches(ctx.margin_bottom_in)
            section.left_margin = Inches(ctx.margin_left_in)
            section.right_margin = Inches(ctx.margin_right_in)

        # Base document styling
        style = doc.styles["Normal"]
        font = style.font
        font.name = ctx.font_family
        font.size = Pt(ctx.font_size_pt)
        font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        style.paragraph_format.line_spacing = ctx.line_spacing

        if ctx.page_numbers:
            for section in doc.sections:
                _add_page_number_header(section, ctx)

        if ctx.title_page:
            _add_title_page(doc, title, ctx.title_page_lines(), ctx)

        # Title at the top
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(title)
        title_run.bold = True
        title_run.font.name = ctx.font_family
        title_run.font.size = Pt(ctx.font_size_pt if ctx.title_page else 18)
        title_para.paragraph_format.space_after = Pt(0 if ctx.title_page else 14)

        if (ctx.author or ctx.institution) and not ctx.title_page:
            meta_p = doc.add_paragraph()
            meta_run = meta_p.add_run(
                f"{ctx.author or ''}{' — ' if ctx.author and ctx.institution else ''}{ctx.institution or ''}"
            )
            meta_run.font.name = ctx.font_family
            meta_run.font.size = Pt(11)
            meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            meta_p.paragraph_format.space_after = Pt(18)

        # Parse text into block chunks
        raw_text = document.text.strip()
        # Split on double newlines
        blocks = re.split(r"\n\s*\n", raw_text)

        in_references = False

        for block in blocks:
            block = block.strip()
            if not block:
                continue

            lines = block.split("\n")
            first_line = lines[0].strip()

            # Skip duplicate title at top if identical
            if first_line.startswith("# ") and first_line[2:].strip().lower() == title.lower():
                # Remainder of block if any
                remaining = "\n".join(lines[1:]).strip()
                if not remaining:
                    continue
                block = remaining
                lines = block.split("\n")
                first_line = lines[0].strip()

            # Check if block is a markdown table
            if len(lines) >= 2 and "|" in first_line and "|" in lines[1]:
                table_data = _parse_markdown_table(lines)
                if table_data:
                    headers, data_rows = table_data
                    tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                    tbl.style = "Table Grid"
                    # Header row
                    hdr_cells = tbl.rows[0].cells
                    for i, h in enumerate(headers):
                        hdr_cells[i].text = h
                        for p in hdr_cells[i].paragraphs:
                            for r in p.runs:
                                r.bold = True
                    # Data rows
                    for r_idx, row in enumerate(data_rows):
                        row_cells = tbl.rows[r_idx + 1].cells
                        for c_idx, val in enumerate(row):
                            row_cells[c_idx].text = val
                    doc.add_paragraph()  # Spacing after table
                    continue

            # Heading 1
            if first_line.startswith("# "):
                h_text = first_line[2:].strip()
                h_p = doc.add_heading(h_text, level=1)
                if ctx.title_page:
                    h_p.paragraph_format.page_break_before = "reference" in h_text.lower()
                    h_p.paragraph_format.space_before = Pt(0)
                    h_p.paragraph_format.space_after = Pt(0)
                    for run in h_p.runs:
                        run.font.name = ctx.font_family
                        run.font.size = Pt(ctx.font_size_pt)
                else:
                    h_p.paragraph_format.space_before = Pt(14)
                    h_p.paragraph_format.space_after = Pt(6)
                    for run in h_p.runs:
                        run.font.name = ctx.font_family
                in_references = "reference" in h_text.lower()
                continue

            # Heading 2
            if first_line.startswith("## "):
                h_text = first_line[3:].strip()
                h_p = doc.add_heading(h_text, level=2)
                if ctx.title_page:
                    h_p.paragraph_format.page_break_before = "reference" in h_text.lower()
                    h_p.paragraph_format.space_before = Pt(0)
                    h_p.paragraph_format.space_after = Pt(0)
                    for run in h_p.runs:
                        run.font.name = ctx.font_family
                        run.font.size = Pt(ctx.font_size_pt)
                else:
                    h_p.paragraph_format.space_before = Pt(12)
                    h_p.paragraph_format.space_after = Pt(4)
                    for run in h_p.runs:
                        run.font.name = ctx.font_family
                in_references = "reference" in h_text.lower()
                continue

            # Heading 3
            if first_line.startswith("### "):
                h_text = first_line[4:].strip()
                h_p = doc.add_heading(h_text, level=3)
                if ctx.title_page:
                    h_p.paragraph_format.space_before = Pt(0)
                    h_p.paragraph_format.space_after = Pt(0)
                    for run in h_p.runs:
                        run.font.name = ctx.font_family
                        run.font.size = Pt(ctx.font_size_pt)
                else:
                    h_p.paragraph_format.space_before = Pt(10)
                    h_p.paragraph_format.space_after = Pt(4)
                    for run in h_p.runs:
                        run.font.name = ctx.font_family
                continue

            # Horizontal rule
            if first_line in ("---", "***", "___"):
                hr_p = doc.add_paragraph()
                hr_p.paragraph_format.space_before = Pt(6)
                hr_p.paragraph_format.space_after = Pt(6)
                run = hr_p.add_run("―" * 40)
                run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
                continue

            # Blockquote
            if first_line.startswith(">"):
                quote_text = "\n".join(l.lstrip("> ").strip() for l in lines)
                q_p = doc.add_paragraph()
                q_p.paragraph_format.left_indent = Inches(0.5)
                q_p.paragraph_format.space_after = Pt(8)
                _add_inline_formatted(q_p, quote_text, ctx)
                for r in q_p.runs:
                    r.italic = True
                continue

            # Bullet list
            if any(l.strip().startswith(("- ", "* ", "• ")) for l in lines):
                for l in lines:
                    l = l.strip()
                    if l.startswith(("- ", "* ", "• ")):
                        bullet_text = l[2:].strip()
                        b_p = doc.add_paragraph(style="List Bullet")
                        b_p.paragraph_format.space_after = Pt(3)
                        _add_inline_formatted(b_p, bullet_text, ctx)
                    else:
                        p = doc.add_paragraph()
                        _add_inline_formatted(p, l, ctx)
                continue

            # Numbered list
            if any(re.match(r"^\d+\.\s+", l.strip()) for l in lines):
                for l in lines:
                    l = l.strip()
                    num_match = re.match(r"^\d+\.\s+(.*)$", l)
                    if num_match:
                        item_text = num_match.group(1)
                        n_p = doc.add_paragraph(style="List Number")
                        n_p.paragraph_format.space_after = Pt(3)
                        _add_inline_formatted(n_p, item_text, ctx)
                    else:
                        p = doc.add_paragraph()
                        _add_inline_formatted(p, l, ctx)
                continue

            # Standard paragraph or reference entry
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = ctx.line_spacing

            if in_references:
                # APA 7 hanging indent for References
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)

            joined_line = " ".join(l.strip() for l in lines)
            _add_inline_formatted(p, joined_line, ctx)

        # AI-Use Disclosure section if requested
        if ctx.include_ai_disclosure and ctx.ai_disclosure_text:
            doc.add_heading("AI-Use Disclosure", level=2)
            disc_p = doc.add_paragraph()
            disc_p.paragraph_format.space_after = Pt(8)
            _add_inline_formatted(disc_p, ctx.ai_disclosure_text.strip(), ctx)

        output_stream = io.BytesIO()
        doc.save(output_stream)
        docx_bytes = output_stream.getvalue()

        return RenderedArtifact(
            format="docx",
            content=docx_bytes,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            title=title,
            filename_extension=".docx",
        )
