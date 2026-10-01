"""Synthetic assignment-package builders. Nothing here is real coursework."""

from __future__ import annotations

from pathlib import Path


def make_pdf(path: Path, text: str) -> None:
    """Write a minimal single-page text PDF that pypdf can extract."""
    esc = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    lines = [esc[i:i + 80] for i in range(0, len(esc), 80)]
    content = "BT /F1 10 Tf 40 760 Td 12 TL " + " T* ".join(f"({ln}) Tj" for ln in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


ZT = (
    "Zero trust architecture removes implicit trust from network location. "
    "Every request is authenticated and authorized using identity and device posture. "
    "Microsegmentation limits lateral movement across enterprise networks. "
)


def build_package(root: Path, *, with_pdf: bool = True) -> Path:
    """instructions + rubric + notes + docx + pdf + pcap + assignment.yaml"""
    import docx

    root.mkdir(parents=True, exist_ok=True)
    (root / "instructions.txt").write_text(
        "Assignment: write a short paper on zero trust architecture.\n"
        "- Explain microsegmentation and lateral movement\n"
        "- Discuss identity and device posture\n"
        "You must cite at least one source.\n"
    )
    (root / "grading-rubric.md").write_text(
        "# Rubric\n1. Defines zero trust architecture clearly\n2. Uses evidence from course material\n"
    )
    (root / "lecture-notes.md").write_text("# Lecture 4\n" + ZT * 4)
    d = docx.Document()
    d.add_heading("Zero Trust Reference", 1)
    d.add_paragraph(ZT * 3)
    d.save(str(root / "reference-document.docx"))
    if with_pdf:
        make_pdf(root / "reference-paper.pdf", ZT * 4)
    (root / "lab-artifact.pcap").write_bytes(b"\xd4\xc3\xb2\xa1" + bytes(256))
    (root / "assignment.yaml").write_text(
        "title: Zero Trust Architecture\n"
        "topic: zero trust architecture microsegmentation\n"
        "target_words: 600\n"
        "source_requirements:\n  minimum_sources: 2\n"
        "materials:\n  directory: .\n"
    )
    return root
